import contextlib
import importlib
import io
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'tools' / 'thesis'))
import tool_cli


class ThesisCliTests(unittest.TestCase):
    def setUp(self):
        base = ROOT / '.cache' / 'test-temp'
        base.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=base, prefix='thesis-cli-')
        self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name)

    def test_default_outputs_are_unique_artifacts(self):
        with patch.object(tool_cli, 'REPOSITORY_ROOT', self.work):
            a = tool_cli.parse_paths('fixture.docx', requires_input=False, argv=[])
            b = tool_cli.parse_paths('fixture.docx', requires_input=False, argv=[])
        self.assertNotEqual(a.output, b.output)
        self.assertTrue(a.output.is_relative_to(self.work / '.artifacts' / 'tool-runs'))

    def test_missing_input_is_rejected(self):
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            tool_cli.parse_paths('fixture.docx', argv=[])

    def test_existing_output_and_source_overwrite_are_rejected(self):
        source = self.work / 'source.docx'
        source.write_bytes(b'synthetic')
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            tool_cli.parse_paths('fixture.docx', argv=['--input', str(source), '--output', str(source)])
        self.assertEqual(source.read_bytes(), b'synthetic')

    def test_all_tools_help_without_reading_input(self):
        for script in (ROOT / 'tools' / 'thesis').glob('*.py'):
            if script.name == 'tool_cli.py':
                continue
            result = subprocess.run([sys.executable, '-B', str(script), '--help'], cwd=self.work,
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, f'{script.name}: {result.stderr}')
            self.assertIn('--output', result.stdout)

    def test_import_remove_tool_has_no_document_side_effect(self):
        module = importlib.import_module('remove_ev_charging_from_proposal')
        self.assertTrue(callable(module.transform))

    def test_synthetic_docx_transforms_without_changing_source(self):
        from docx import Document
        source = self.work / 'source with spaces.docx'
        output = self.work / 'result.docx'
        doc = Document()
        doc.add_paragraph('Synthetic content preserved')
        doc.add_table(rows=1, cols=1)
        doc.add_table(rows=11, cols=2)
        doc.save(source)
        original = source.read_bytes()
        script = ROOT / 'tools' / 'thesis' / 'remove_ev_charging_from_proposal.py'
        result = subprocess.run([sys.executable, '-B', str(script), '--input', str(source),
                                 '--output', str(output)], cwd=self.work, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(Document(output).paragraphs[0].text, 'Synthetic content preserved')

    def test_transform_refuses_output_created_after_cli_validation(self):
        from docx import Document
        from remove_ev_charging_from_proposal import transform
        source = self.work / 'source.docx'
        output = self.work / 'already-created.docx'
        doc = Document()
        doc.add_table(rows=1, cols=1)
        doc.add_table(rows=11, cols=2)
        doc.save(source)
        output.write_bytes(b'other writer owns this')
        with self.assertRaises(FileExistsError):
            transform(source, output)
        self.assertEqual(output.read_bytes(), b'other writer owns this')

    def test_template_hash_guard_is_preserved(self):
        from fill_thesis_proposal_template import EXPECTED_SOURCE_SHA256
        source = self.work / 'not-the-original.docx'
        source.write_bytes(b'synthetic not original template')
        result = subprocess.run([sys.executable, '-B', str(ROOT / 'tools/thesis/fill_thesis_proposal_template.py'),
                                 '--input', str(source), '--output', str(self.work / 'rejected.docx')],
                                capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('Reference changed', result.stderr)
        self.assertEqual(len(EXPECTED_SOURCE_SHA256), 64)
        self.assertFalse((self.work / 'rejected.docx').exists())


if __name__ == '__main__':
    import argparse
    import xml.etree.ElementTree as ET
    parser = argparse.ArgumentParser()
    parser.add_argument('--junitxml', type=Path)
    args = parser.parse_args()
    cases = list(unittest.defaultTestLoader.loadTestsFromTestCase(ThesisCliTests))
    result = unittest.TextTestRunner(verbosity=2).run(unittest.TestSuite(cases))
    if args.junitxml:
        report = ET.Element('testsuite', tests=str(result.testsRun), failures=str(len(result.failures)),
                            errors=str(len(result.errors)), skipped=str(len(result.skipped)))
        failures = {test.id(): message for test, message in result.failures}
        errors = {test.id(): message for test, message in result.errors}
        skipped = {test.id(): message for test, message in result.skipped}
        for case in cases:
            item = ET.SubElement(report, 'testcase', name=case.id())
            for tag, outcomes in [('failure', failures), ('error', errors), ('skipped', skipped)]:
                if case.id() in outcomes:
                    ET.SubElement(item, tag).text = outcomes[case.id()]
        ET.ElementTree(report).write(args.junitxml, encoding='utf-8', xml_declaration=True)
    sys.exit(0 if result.wasSuccessful() else 1)
