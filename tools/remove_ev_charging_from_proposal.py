from pathlib import Path
from docx import Document


SOURCE = Path(r"D:\Code\KhoaLuan\De_cuong_chi_tiet_KLCN_QL_BaiDoXe.docx")
OUTPUT = Path(r"D:\Code\KhoaLuan\De_cuong_chi_tiet_KLCN_QL_BaiDoXe_KHONG_SAC_XE_DIEN.docx")


def replace_after_label(paragraph, new_body: str) -> None:
    """Preserve the bold label run and replace only the descriptive body."""
    if len(paragraph.runs) < 2:
        raise ValueError(f"Unexpected run structure: {paragraph.text!r}")
    paragraph.runs[1].text = " " + new_body
    for run in paragraph.runs[2:]:
        run.text = ""


def replace_entire_paragraph(paragraph, new_text: str) -> None:
    """Keep paragraph/list formatting and the formatting of its first run."""
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


doc = Document(SOURCE)

# Inline edits outside NV06. Vehicle-type classification remains in scope;
# only charging-area, charging-right, and charging-session management is removed.
body_replacements = {
    "Đối tượng khảo sát:": "Ban quản lý, nhân viên trạm gác, cư dân/chủ căn hộ, thành viên hộ gia đình và khách vãng lai.",
    "Phạm vi quan sát:": "Cổng vào, cổng ra, quầy đăng ký thẻ, khu đỗ xe máy, khu ô tô, khu xe điện và khu xử lý sự cố.",
    "Nhóm bảng giao dịch:": "parking_sessions, recognition_results, media_files, payments, alerts và audit_logs.",
    "Dashboard vận hành:": "Số xe trong bãi, sức chứa, lượt vào-ra, lượt OPEN lâu, cảnh báo chưa xử lý và trạng thái API/AI/camera.",
}

for paragraph in doc.paragraphs:
    for label, body in body_replacements.items():
        if paragraph.text.startswith(label):
            replace_after_label(paragraph, body)
            break

# Only the registration Bước 9 mentions charging. Other Bước 9 paragraphs
# belong to the entry/exit workflows and must remain unchanged.
for paragraph in doc.paragraphs:
    if paragraph.text == "Bước 9: Cấp thẻ, liên kết xe và lựa chọn gói 30 ngày phù hợp; hiển thị ngày bắt đầu-hết hạn và quyền sạc.":
        replace_after_label(paragraph, "Cấp thẻ, liên kết xe và lựa chọn gói 30 ngày phù hợp; hiển thị rõ ngày bắt đầu và ngày hết hạn.")

# Rewrite the remaining NV06 content as pure vehicle classification.
for paragraph in doc.paragraphs:
    text = paragraph.text
    if text == "Nghiệp vụ NV06 - Phân loại phương tiện và quản lý quyền sạc xe điện":
        replace_entire_paragraph(paragraph, "Nghiệp vụ NV06 - Phân loại phương tiện")
    elif text.startswith("Mục đích: Xác định nhóm xe để áp dụng quy định/giá"):
        replace_after_label(paragraph, "Xác định đúng nhóm phương tiện để áp dụng quy định, giới hạn đăng ký và bảng giá phù hợp.")
    elif text.startswith("Dữ liệu lưu: Dự đoán, độ tin cậy, loại chính thức"):
        replace_after_label(paragraph, "Dự đoán AI, độ tin cậy, loại phương tiện chính thức, thông tin kỹ thuật và người xác minh.")
    elif text.startswith("Kết quả: Áp đúng nhóm giá"):
        replace_after_label(paragraph, "Phương tiện được gán đúng nhóm quản lý và PriceRule; sai lệch giữa AI với hồ sơ được cảnh báo.")
    elif text.startswith("Tiêu chí chấp nhận: Mọi quyết định khác AI"):
        replace_after_label(paragraph, "Mọi trường hợp nhân viên xác minh khác dự đoán AI đều có lý do và lịch sử truy vết.")
    elif text.startswith("Nghiệp vụ NV08 - Tra cứu, cảnh báo"):
        # Avoid placing the NV08 heading and its first detail line on the same
        # baseline at the bottom of a nearly full page in Microsoft Word.
        paragraph.paragraph_format.page_break_before = True

# Remove charging-only workflow paragraphs from NV06.
charging_only_prefixes = (
    "Bước đăng ký 3: Nếu xe điện, lưu quyền sạc",
    "Bước sạc 1:",
    "Bước sạc 2:",
    "Bước sạc 3:",
    "Bước sạc 4:",
    "Ngoại lệ 1: Xe xăng vào khu sạc",
)
for paragraph in list(doc.paragraphs):
    if paragraph.text.startswith(charging_only_prefixes):
        delete_paragraph(paragraph)

# Update week 10 without changing table structure or paragraph formatting.
schedule_cell = doc.tables[1].rows[10].cells[1]
for paragraph in schedule_cell.paragraphs:
    if paragraph.text.startswith("Công việc:"):
        replace_after_label(paragraph, "Hoàn thiện NV06-NV08: phân loại phương tiện, cảnh báo, tra cứu, doanh thu và dashboard.")
    elif paragraph.text.startswith("Sản phẩm:"):
        replace_after_label(paragraph, "Báo cáo đối soát; cảnh báo; kết quả phân loại phương tiện và dashboard.")

doc.save(OUTPUT)
print(OUTPUT)
