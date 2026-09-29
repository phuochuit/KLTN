from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path


OUT = Path(r"D:\Code\KhoaLuan\Ke_hoach_do_an_bai_giu_xe_ANPR_Spring_Boot.docx")

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"
GREEN = "2F6B4F"
GOLD = "8A6400"
RED = "9B1C1C"
BLACK = "111827"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            w = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(w))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(w / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Calibri", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)


def set_repeat_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("ĐỒ ÁN HỆ THỐNG BÃI GIỮ XE ANPR  |  KẾ HOẠCH TRIỂN KHAI")
    set_run_font(r, size=8.5, bold=True, color=MID_GRAY)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("Trang ")
    set_run_font(r, size=8.5, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r._r.append(fld_char1)
    r._r.append(instr)
    r._r.append(fld_char2)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    set_repeat_header_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Title", 28, NAVY, 0, 8),
        ("Subtitle", 14, MID_GRAY, 0, 12),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = rgb(color)
        st.font.bold = name != "Subtitle"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.375)
        st.paragraph_format.first_line_indent = Inches(-0.188)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Callout"]
    callout.font.name = "Calibri"
    callout.font.size = Pt(10.5)
    callout.font.color.rgb = rgb(NAVY)
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.18)
    callout.paragraph_format.space_before = Pt(6)
    callout.paragraph_format.space_after = Pt(10)
    callout.paragraph_format.line_spacing = 1.2
    return doc


def add_shading_to_paragraph(paragraph, fill=LIGHT_BLUE, border=BLUE):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border)
    pbdr.append(left)
    ppr.append(pbdr)


def add_callout(doc, label, text, fill=LIGHT_BLUE, border=BLUE):
    p = doc.add_paragraph(style="Callout")
    add_shading_to_paragraph(p, fill, border)
    r = p.add_run(label + ": ")
    set_run_font(r, size=10.5, bold=True, color=NAVY)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=NAVY)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_table(doc, headers, rows, widths_dxa, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, size=font_size, bold=True, color=NAVY)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if ridx % 2 == 1:
                set_cell_shading(cells[i], "FAFBFC")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=BLACK)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_break(doc):
    doc.add_page_break()


def build_document():
    doc = setup_document()

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("KẾ HOẠCH TRIỂN KHAI ĐỒ ÁN")
    set_run_font(r, size=12, bold=True, color=BLUE)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HỆ THỐNG QUẢN LÝ BÃI GIỮ XE\nTÍCH HỢP NHẬN DIỆN BIỂN SỐ ANPR")
    set_run_font(r, size=27, bold=True, color=NAVY)

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Web Java Spring Boot • Desktop trạm gác • YOLO + OCR")
    set_run_font(r, size=14, color=MID_GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    r = p.add_run("Tài liệu phân tích, thiết kế, phân công và lộ trình thực hiện cho nhóm 3 người")
    set_run_font(r, size=11.5, italic=True, color=DARK_BLUE)

    add_callout(doc, "Định hướng công nghệ", "Web được xây dựng bằng Java Spring Boot (Spring MVC + Thymeleaf). Spring Boot cung cấp REST API cho ứng dụng Desktop. Dịch vụ ANPR sử dụng Python để tận dụng YOLO và OCR.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run("NHÓM THỰC HIỆN: 03 THÀNH VIÊN")
    set_run_font(r, size=10.5, bold=True, color=MID_GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Tháng 08/2026")
    set_run_font(r, size=10.5, color=MID_GRAY)

    add_page_break(doc)

    # Document use and contents
    doc.add_heading("HƯỚNG DẪN SỬ DỤNG TÀI LIỆU", level=1)
    doc.add_paragraph("Tài liệu này là kế hoạch làm việc thực tế cho nhóm ba người. Nhóm có thể dùng để chốt phạm vi với giảng viên, tạo backlog, phân công nhiệm vụ, thiết kế cơ sở dữ liệu và chuẩn bị nội dung báo cáo đồ án.")
    add_callout(doc, "Nguyên tắc ưu tiên", "Hoàn thành nghiệp vụ xe vào/ra và tính phí trước; sau đó cải thiện độ chính xác ANPR, báo cáo và giao diện. Không mở rộng tính năng khi luồng cốt lõi chưa ổn định.", fill="FFF7E6", border=GOLD)

    doc.add_heading("MỤC LỤC NỘI DUNG", level=1)
    toc = [
        "1. Tóm tắt đề tài và phạm vi", "2. Kiến trúc hệ thống và công nghệ", "3. Yêu cầu chức năng và phân quyền",
        "4. Phân tích nghiệp vụ", "5. Thiết kế dữ liệu", "6. Thiết kế Web Java Spring Boot", "7. Thiết kế Desktop",
        "8. Thiết kế ANPR", "9. Phân chia công việc cho ba thành viên", "10. Kế hoạch 14 tuần", "11. Kiểm thử và tiêu chí nghiệm thu",
        "12. Rủi ro và phương án xử lý", "13. Hướng nâng cấp sau khi hoàn thành", "14. Checklist bảo vệ đồ án"
    ]
    add_numbers(doc, toc)

    # Section 1
    doc.add_heading("1. TÓM TẮT ĐỀ TÀI VÀ PHẠM VI", level=1)
    doc.add_heading("1.1. Mục tiêu", level=2)
    doc.add_paragraph("Xây dựng hệ thống tự động hóa quy trình kiểm soát xe ra vào tại tòa nhà hoặc trường học, giảm thời gian chờ tại cổng, hạn chế ghi nhận thủ công và tăng khả năng truy vết sự cố.")
    add_bullets(doc, [
        "Nhận hình ảnh từ camera cổng vào và cổng ra.",
        "Dùng YOLO phát hiện, khoanh vùng biển số; dùng OCR nhận dạng ký tự.",
        "Đối chiếu biển số với xe/thẻ đã đăng ký và lượt xe đang mở.",
        "Tự động ghi nhận thời gian vào/ra, tính phí và cảnh báo ngoại lệ.",
        "Cung cấp Web quản trị để quản lý xe, thẻ, bảng giá, lịch sử và doanh thu."
    ])
    doc.add_heading("1.2. Phạm vi MVP bắt buộc", level=2)
    add_table(doc, ["Nhóm nghiệp vụ", "Kết quả cần đạt"], [
        ("Quản lý xe/thẻ", "CRUD xe và chủ xe; cấp, khóa, gia hạn thẻ; tra cứu theo biển số hoặc mã thẻ."),
        ("Kiểm soát vào/ra", "Nhận diện biển số; đối chiếu; ghi ảnh, thời gian, nhân viên và trạng thái lượt gửi."),
        ("Tính phí", "Áp dụng bảng giá theo loại xe/vé; tạo giao dịch khi xe ra; không bị ảnh hưởng bởi thay đổi giá về sau."),
        ("Báo cáo", "Doanh thu theo ngày/tháng; số lượt; xe đang trong bãi; lịch sử và ảnh bằng chứng."),
        ("Phân quyền", "Hai vai trò: Quản trị viên/Ban quản lý và Nhân viên trạm gác.")
    ], [2300, 7060], 9.5)
    doc.add_heading("1.3. Ngoài phạm vi bản đầu", level=2)
    add_bullets(doc, ["Điều khiển barie thật", "Thanh toán ngân hàng/ví điện tử", "Ứng dụng mobile", "Nhiều chi nhánh", "Nhận diện khuôn mặt", "Tự huấn luyện OCR từ đầu"])

    # Section 2
    doc.add_heading("2. KIẾN TRÚC HỆ THỐNG VÀ CÔNG NGHỆ", level=1)
    doc.add_heading("2.1. Kiến trúc logic", level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Camera → Desktop WPF → ANPR Python\nDesktop/Web → Spring Boot → PostgreSQL\nSpring Boot → lưu ảnh, đối chiếu, tính phí, báo cáo")
    set_run_font(r, size=11.5, bold=True, color=NAVY)
    add_callout(doc, "Quyết định kiến trúc", "Spring Boot là nơi duy nhất xử lý nghiệp vụ và truy cập cơ sở dữ liệu. Desktop và Web không tự tính phí. ANPR chỉ nhận ảnh và trả kết quả nhận dạng, không truy cập database.")
    doc.add_heading("2.2. Stack công nghệ đề xuất", level=2)
    add_table(doc, ["Thành phần", "Công nghệ", "Vai trò"], [
        ("Web + Backend", "Java 21, Spring Boot 3, Spring MVC, Thymeleaf", "Giao diện quản trị, REST API và nghiệp vụ trung tâm."),
        ("Bảo mật", "Spring Security, JWT hoặc session", "Đăng nhập, phân quyền theo vai trò."),
        ("Dữ liệu", "Spring Data JPA, PostgreSQL", "ORM, migration và lưu dữ liệu nghiệp vụ."),
        ("Desktop", "C# .NET 8, WPF, OpenCvSharp", "Kết nối camera, hiển thị và xác nhận xe vào/ra."),
        ("ANPR", "Python, FastAPI, Ultralytics YOLO, PaddleOCR/EasyOCR", "Phát hiện biển số và nhận dạng ký tự."),
        ("Kiểm thử", "JUnit 5, Mockito, Testcontainers, Postman", "Kiểm thử đơn vị, tích hợp và API."),
        ("Triển khai", "Docker Compose", "Khởi chạy Spring Boot, PostgreSQL và ANPR nhất quán.")
    ], [1800, 3000, 4560], 8.8)

    # Section 3
    doc.add_heading("3. YÊU CẦU CHỨC NĂNG VÀ PHÂN QUYỀN", level=1)
    doc.add_heading("3.1. Web quản trị - Java Spring Boot", level=2)
    add_bullets(doc, [
        "Đăng nhập, đăng xuất và đổi mật khẩu.", "Dashboard: tổng doanh thu, số lượt, số xe đang trong bãi và số cảnh báo.",
        "Quản lý tài khoản, xe, chủ xe, thẻ, loại xe và bảng giá.", "Tra cứu lịch sử vào/ra theo thời gian, biển số, thẻ, trạng thái.",
        "Xem ảnh toàn cảnh, ảnh biển số và kết quả AI.", "Xuất báo cáo Excel/PDF nếu còn thời gian."
    ])
    doc.add_heading("3.2. Desktop trạm gác", level=2)
    add_bullets(doc, [
        "Kết nối webcam/IP camera và hiển thị video trực tiếp.", "Chụp khung hình tự động hoặc bằng nút bấm.",
        "Gửi ảnh đến ANPR, hiển thị bounding box, biển số và confidence.", "Cho phép sửa biển số trước khi xác nhận.",
        "Gọi API xe vào, xem trước xe ra, xác nhận thanh toán và đóng lượt.", "Cảnh báo màu/âm thanh khi xe lạ, thẻ bị khóa hoặc biển số không khớp."
    ])
    doc.add_heading("3.3. Ma trận phân quyền", level=2)
    add_table(doc, ["Chức năng", "Quản trị viên", "Nhân viên trạm gác"], [
        ("Quản lý tài khoản và vai trò", "Được phép", "Không"),
        ("Quản lý xe, thẻ, bảng giá", "Được phép", "Chỉ tra cứu"),
        ("Xem báo cáo doanh thu", "Đầy đủ", "Không hoặc giới hạn ca trực"),
        ("Ghi nhận xe vào/ra", "Được phép", "Được phép"),
        ("Sửa kết quả nhận dạng", "Được phép", "Được phép, có audit log"),
        ("Xem nhật ký hệ thống", "Được phép", "Không")
    ], [4300, 2300, 2760], 9)

    # Section 4
    doc.add_heading("4. PHÂN TÍCH NGHIỆP VỤ", level=1)
    doc.add_heading("4.1. Luồng xe vào", level=2)
    add_numbers(doc, [
        "Camera chụp ảnh xe tại vùng dừng.", "Desktop gửi ảnh sang dịch vụ ANPR.", "YOLO phát hiện vùng biển số; OCR nhận dạng ký tự.",
        "Desktop gửi biển số, confidence và ảnh sang Spring Boot.", "Spring Boot chuẩn hóa biển số, tra cứu xe/thẻ và kiểm tra lượt đang mở.",
        "Nhân viên xem cảnh báo, sửa biển số nếu cần và xác nhận.", "Spring Boot tạo ParkingSession trạng thái OPEN và ghi audit log."
    ])
    doc.add_heading("4.2. Luồng xe ra", level=2)
    add_numbers(doc, [
        "ANPR nhận dạng biển số tại cổng ra.", "Desktop gọi API exit-preview để tìm lượt OPEN và tính thử phí.",
        "Spring Boot so sánh biển số vào/ra, trạng thái thẻ và loại vé.", "Nhân viên xử lý cảnh báo và xác nhận thanh toán.",
        "Desktop gọi exit-confirm kèm mã yêu cầu duy nhất.", "Spring Boot đóng lượt, tạo Payment và lưu ảnh ra."
    ])
    doc.add_heading("4.3. Quy tắc tính phí MVP", level=2)
    add_bullets(doc, [
        "Xe tháng còn hạn: phí lượt bằng 0.", "Xe lượt cùng ngày: áp dụng giá lượt theo loại xe.",
        "Xe qua ngày: giá lượt cộng số đêm nhân phí qua đêm.", "Bảng giá có ngày bắt đầu áp dụng; không sửa trực tiếp bản giá cũ.",
        "Payment lưu số tiền cuối cùng và pricingRuleId đã áp dụng để bảo toàn lịch sử."
    ])
    doc.add_heading("4.4. Ngoại lệ phải xử lý", level=2)
    add_table(doc, ["Tình huống", "Cách xử lý đề xuất"], [
        ("OCR confidence thấp", "Yêu cầu nhân viên nhập/sửa biển số; không tự động xác nhận."),
        ("Không tìm thấy lượt vào", "Cảnh báo đỏ; cho quản trị viên hoặc nhân viên có quyền tạo xử lý thủ công."),
        ("Biển số vào/ra khác", "Hiển thị hai ảnh và gợi ý biển số gần giống; bắt buộc xác nhận."),
        ("Thẻ bị khóa/mất", "Không tự động cho qua; ghi Alert và audit log."),
        ("Bấm xác nhận nhiều lần", "Dùng idempotencyKey để API chỉ tạo một giao dịch."),
        ("ANPR mất kết nối", "Cho phép nhập biển số thủ công và lưu nguyên nhân xử lý thủ công.")
    ], [2750, 6610], 9)

    # Section 5
    doc.add_heading("5. THIẾT KẾ DỮ LIỆU", level=1)
    doc.add_heading("5.1. Các bảng chính", level=2)
    add_table(doc, ["Bảng", "Mục đích", "Trường quan trọng"], [
        ("users / roles", "Tài khoản và quyền", "username, passwordHash, role, status"),
        ("vehicle_owners", "Chủ xe", "name, phone, studentCode/employeeCode"),
        ("vehicles", "Xe đăng ký", "plateNormalized, vehicleTypeId, ownerId, status"),
        ("parking_cards", "Thẻ/vé", "cardCode, type, startDate, expiryDate, status"),
        ("pricing_rules", "Bảng giá", "vehicleTypeId, ticketType, basePrice, overnightFee, effectiveFrom"),
        ("parking_sessions", "Lượt gửi xe", "entry/exit plate, time, image, operator, status, fee"),
        ("recognition_results", "Kết quả AI", "rawText, normalizedText, detectConf, ocrConf, bbox, modelVersion"),
        ("payments", "Giao dịch", "sessionId, amount, method, paidAt, status"),
        ("alerts / audit_logs", "Cảnh báo và truy vết", "type, severity, actor, action, before/after data")
    ], [2000, 2600, 4760], 8.6)
    doc.add_heading("5.2. Ràng buộc quan trọng", level=2)
    add_bullets(doc, [
        "plateNormalized có index và không chứa khoảng trắng, dấu chấm hoặc gạch ngang.",
        "Một xe chỉ có tối đa một ParkingSession trạng thái OPEN.", "Một ParkingSession chỉ có tối đa một Payment thành công.",
        "Không xóa vật lý bản ghi đã phát sinh lịch sử; dùng trạng thái INACTIVE/DELETED.",
        "Thời gian lưu theo UTC trong database; hiển thị theo múi giờ Asia/Bangkok.", "Mật khẩu luôn được băm bằng BCrypt/Argon2."
    ])

    # Section 6
    doc.add_heading("6. THIẾT KẾ WEB JAVA SPRING BOOT", level=1)
    doc.add_heading("6.1. Cấu trúc module", level=2)
    p = doc.add_paragraph()
    r = p.add_run("src/main/java/com/example/parking/")
    set_run_font(r, name="Consolas", size=10.5, bold=True, color=NAVY)
    modules = [
        "config - Security, CORS, Jackson, OpenAPI", "auth - đăng nhập và phân quyền", "vehicle - xe, chủ xe, loại xe",
        "card - thẻ và gán thẻ cho xe", "pricing - bảng giá và thuật toán tính phí", "session - lượt xe vào/ra",
        "payment - giao dịch", "report - dashboard và báo cáo", "alert - cảnh báo", "audit - nhật ký thao tác",
        "storage - lưu ảnh local hoặc object storage", "common - exception, response, validation, utilities"
    ]
    add_bullets(doc, modules)
    doc.add_heading("6.2. Layer chuẩn", level=2)
    add_table(doc, ["Layer", "Trách nhiệm"], [
        ("Controller", "Nhận HTTP request, validate đầu vào, trả view Thymeleaf hoặc JSON."),
        ("Service", "Chứa nghiệp vụ, transaction, đối chiếu xe/thẻ và tính phí."),
        ("Repository", "Truy vấn database bằng Spring Data JPA; không chứa nghiệp vụ."),
        ("Entity", "Ánh xạ bảng; hạn chế trả trực tiếp ra API."),
        ("DTO", "Request/response riêng; dùng Bean Validation."),
        ("Mapper", "Chuyển Entity ↔ DTO, có thể dùng MapStruct.")
    ], [1900, 7460], 9.3)
    doc.add_heading("6.3. API cốt lõi", level=2)
    add_table(doc, ["Method", "Endpoint", "Mục đích"], [
        ("POST", "/api/auth/login", "Đăng nhập Desktop/API"),
        ("GET/POST/PUT", "/api/vehicles", "Quản lý xe"),
        ("GET/POST/PUT", "/api/cards", "Quản lý thẻ"),
        ("GET/POST/PUT", "/api/pricing-rules", "Quản lý bảng giá"),
        ("POST", "/api/parking-sessions/entry", "Tạo lượt xe vào"),
        ("POST", "/api/parking-sessions/exit-preview", "Tìm lượt và tính phí dự kiến"),
        ("POST", "/api/parking-sessions/exit-confirm", "Đóng lượt và tạo thanh toán"),
        ("GET", "/api/reports/dashboard", "Số liệu dashboard"),
        ("GET", "/api/reports/revenue", "Doanh thu theo khoảng thời gian")
    ], [1100, 3750, 4510], 8.8)
    doc.add_heading("6.4. Màn hình Thymeleaf", level=2)
    add_bullets(doc, [
        "login.html", "dashboard.html", "vehicles/list.html và form.html", "cards/list.html và form.html",
        "pricing/list.html và form.html", "sessions/history.html và detail.html", "reports/revenue.html", "users/list.html"
    ])
    add_callout(doc, "Lưu ý", "Nếu nhóm muốn giao diện hiện đại hơn, có thể thay Thymeleaf bằng React ở giai đoạn nâng cấp. Với nhóm ba người và thời gian hạn chế, Thymeleaf giúp giảm một hệ sinh thái phải học và triển khai.", fill="FFF7E6", border=GOLD)

    # Section 7
    doc.add_heading("7. THIẾT KẾ DESKTOP TRẠM GÁC", level=1)
    doc.add_heading("7.1. Các màn hình", level=2)
    add_bullets(doc, [
        "Đăng nhập", "Màn hình vận hành cổng vào", "Màn hình vận hành cổng ra", "Cấu hình camera/API",
        "Danh sách lượt đang mở", "Lịch sử thao tác trong ca"
    ])
    doc.add_heading("7.2. Thành phần màn hình vận hành", level=2)
    add_table(doc, ["Vùng", "Nội dung"], [
        ("Camera", "Video trực tiếp, trạng thái kết nối, nút chụp/nhận dạng."),
        ("Kết quả ANPR", "Ảnh crop, biển số, confidence YOLO/OCR, ô sửa biển số."),
        ("Đối chiếu", "Thông tin xe, chủ xe, thẻ, lượt vào và cảnh báo."),
        ("Thao tác", "Xác nhận vào; xem trước phí; xác nhận ra; hủy; nhập thủ công."),
        ("Trạng thái hệ thống", "Camera, ANPR service, Spring Boot API và mạng.")
    ], [2200, 7160], 9.4)
    doc.add_heading("7.3. Yêu cầu kỹ thuật", level=2)
    add_bullets(doc, [
        "Không khóa UI khi gọi camera/API; dùng async/await và CancellationToken.", "Giới hạn tần suất nhận dạng để tránh gửi mọi frame.",
        "Có timeout, retry có kiểm soát và thông báo lỗi dễ hiểu.", "Không lưu mật khẩu dạng rõ; token đặt trong storage an toàn.",
        "Có chế độ chọn ảnh/video để demo khi không có camera thật."
    ])

    # Section 8
    doc.add_heading("8. THIẾT KẾ ANPR: YOLO + OCR", level=1)
    doc.add_heading("8.1. Pipeline", level=2)
    add_numbers(doc, [
        "Nhận ảnh từ Desktop qua POST /recognize.", "Resize/letterbox ảnh theo đầu vào của YOLO.", "YOLO trả bounding box và detection confidence.",
        "Crop biển số; căn chỉnh phối cảnh nếu cần.", "Tiền xử lý nhiều biến thể: grayscale, tăng tương phản, sharpen/threshold có điều kiện.",
        "OCR đọc ký tự và trả confidence.", "Chuẩn hóa chuỗi, kiểm tra cấu trúc biển số Việt Nam và trả JSON."
    ])
    doc.add_heading("8.2. Chuẩn hóa biển số", level=2)
    add_bullets(doc, [
        "Chuyển thành chữ hoa; bỏ khoảng trắng, dấu chấm và gạch ngang.", "Chỉ giữ ký tự A-Z và 0-9.",
        "Sửa nhầm O/0, I/1, B/8 theo vị trí và mẫu biển số; không thay thế mù quáng.",
        "Cho phép đối chiếu gần đúng bằng Levenshtein để gợi ý, nhưng không tự mở cổng nếu chưa khớp chính xác."
    ])
    doc.add_heading("8.3. API trả về", level=2)
    p = doc.add_paragraph()
    r = p.add_run('{\n  "plateText": "59A12345",\n  "detectionConfidence": 0.94,\n  "ocrConfidence": 0.87,\n  "boundingBox": {"x": 320, "y": 185, "width": 170, "height": 64},\n  "modelVersion": "yolo-plate-v1"\n}')
    set_run_font(r, name="Consolas", size=9.3, color=BLACK)
    p.paragraph_format.left_indent = Inches(0.25)
    add_shading_to_paragraph(p, "F7F8FA", MID_GRAY)
    doc.add_heading("8.4. Đánh giá mô hình", level=2)
    add_table(doc, ["Chỉ số", "Cách đo"], [
        ("mAP/Precision/Recall", "Đánh giá bước YOLO trên tập test tách riêng."),
        ("Độ chính xác chuỗi", "Tỷ lệ biển số nhận đúng toàn bộ ký tự."),
        ("Độ chính xác ký tự", "Số ký tự đúng trên tổng số ký tự."),
        ("Thời gian xử lý", "Trung bình và p95 cho một ảnh."),
        ("Tỷ lệ sửa thủ công", "Số lượt nhân viên sửa kết quả trên tổng lượt nhận dạng.")
    ], [2600, 6760], 9.3)

    # Section 9
    doc.add_heading("9. PHÂN CHIA CÔNG VIỆC CHO NHÓM 3 NGƯỜI", level=1)
    add_table(doc, ["Thành viên", "Trách nhiệm chính", "Đầu ra bắt buộc"], [
        ("TV1 - Web/Backend", "Spring Boot, Thymeleaf, PostgreSQL, Security, nghiệp vụ và báo cáo.", "ERD, migration, API, Web quản trị, test nghiệp vụ."),
        ("TV2 - Desktop", "WPF, camera, tích hợp ANPR/API, giao diện vận hành và xử lý lỗi.", "Desktop vào/ra, camera/video demo, cảnh báo, log client."),
        ("TV3 - ANPR + hỗ trợ Web", "Dataset, YOLO, OCR, FastAPI, đánh giá; hỗ trợ giao diện Thymeleaf.", "ANPR service, model, tập test, báo cáo chỉ số, Dockerfile.")
    ], [1800, 4100, 3460], 8.7)
    doc.add_heading("9.1. Việc làm chung", level=2)
    add_bullets(doc, [
        "Cùng chốt use case, ERD và hợp đồng API trước khi code.", "Mỗi pull request phải có ít nhất một thành viên khác review.",
        "Mỗi tuần tích hợp ít nhất một lần; không đợi đến cuối mới ghép hệ thống.", "Mỗi phần phải có một người dự phòng biết cách chạy và xử lý lỗi cơ bản.",
        "Cả nhóm cùng chuẩn bị dữ liệu demo, tài liệu, slide và luyện bảo vệ."
    ])
    doc.add_heading("9.2. Cấu trúc repository", level=2)
    add_bullets(doc, ["/spring-web", "/desktop-wpf", "/anpr-service", "/database", "/docs", "/deployment"])

    # Section 10
    doc.add_heading("10. KẾ HOẠCH THỰC HIỆN 14 TUẦN", level=1)
    rows = [
        ("1-2", "Phân tích & thiết kế", "Use case, luồng vào/ra, ERD, mockup, API contract", "Tài liệu baseline được cả nhóm duyệt"),
        ("3-4", "Nền tảng", "Spring Security, CRUD xe/thẻ; Desktop mở camera; ANPR ảnh tĩnh", "CRUD xuyên suốt và nhận dạng thử"),
        ("5-6", "Xe vào", "Entry API, lưu ảnh, đối chiếu xe/thẻ, Desktop xác nhận", "Demo hoàn chỉnh luồng vào"),
        ("7-8", "Xe ra & tính phí", "Exit preview/confirm, bảng giá, payment, cảnh báo", "Demo vào → ra → thanh toán"),
        ("9-10", "Báo cáo & AI", "Dashboard, doanh thu, fine-tune/tối ưu OCR, tập test", "Đủ ba nghiệp vụ bắt buộc"),
        ("11-12", "Kiểm thử", "Phân quyền, lỗi mạng, đồng thời, tính phí, bảo mật", "Báo cáo test và danh sách lỗi đã sửa"),
        ("13-14", "Hoàn thiện", "Triển khai, tài liệu, video, slide, luyện bảo vệ", "Release demo ổn định")
    ]
    add_table(doc, ["Tuần", "Giai đoạn", "Công việc", "Mốc nghiệm thu"], rows, [900, 1800, 4100, 2560], 8.3)
    doc.add_heading("10.1. Definition of Done cho mỗi chức năng", level=2)
    add_bullets(doc, [
        "Có tiêu chí chấp nhận rõ ràng.", "Code đã được review và merge vào develop.", "Có test phù hợp và không làm hỏng chức năng cũ.",
        "Đã chạy tích hợp với ít nhất một thành phần liên quan.", "Không chứa mật khẩu, token hoặc file model lớn trong Git.", "Đã cập nhật hướng dẫn chạy nếu cấu hình thay đổi."
    ])

    # Section 11
    doc.add_heading("11. KIỂM THỬ VÀ TIÊU CHÍ NGHIỆM THU", level=1)
    doc.add_heading("11.1. Nhóm kiểm thử", level=2)
    add_table(doc, ["Loại kiểm thử", "Nội dung"], [
        ("Unit test", "PricingService, chuẩn hóa biển số, phân quyền, validation."),
        ("Integration test", "Repository/PostgreSQL, entry/exit transaction, idempotency."),
        ("API test", "Mã trạng thái, dữ liệu sai, token hết hạn, quyền truy cập."),
        ("Desktop test", "Camera mất kết nối, timeout API, thao tác lặp, nhập thủ công."),
        ("ANPR test", "Ngày/đêm, nghiêng, mờ, xe máy/ô tô, biển số bẩn."),
        ("End-to-end", "Xe đăng ký và xe lạ đi qua trọn luồng vào → ra → báo cáo.")
    ], [2300, 7060], 9.3)
    doc.add_heading("11.2. Checklist nghiệm thu MVP", level=2)
    checks = [
        "Camera/video hiển thị được trên Desktop.", "YOLO khoanh vùng biển số và OCR trả ký tự/confidence.", "Nhân viên sửa được kết quả trước khi xác nhận.",
        "Tạo được lượt xe vào và lưu ảnh bằng chứng.", "Xe ra tìm đúng lượt vào; cảnh báo khi không khớp.", "Tính phí và tạo Payment chính xác.",
        "Web quản lý được xe, thẻ, bảng giá.", "Tra cứu được lịch sử và ảnh.", "Dashboard hiển thị doanh thu ngày/tháng.",
        "Hai vai trò bị giới hạn quyền đúng.", "Thao tác quan trọng có audit log.", "Có số liệu đánh giá ANPR trên tập test riêng."
    ]
    for item in checks:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run("☐ " + item)
        set_run_font(r, size=10.8, color=BLACK)

    # Section 12
    doc.add_heading("12. RỦI RO VÀ PHƯƠNG ÁN XỬ LÝ", level=1)
    add_table(doc, ["Rủi ro", "Dấu hiệu", "Giảm thiểu"], [
        ("ANPR thiếu ổn định", "Sai khi tối, nghiêng, chuyển động", "Cố định vùng dừng, chụp nhiều frame, confidence threshold, sửa thủ công."),
        ("Phạm vi quá lớn", "Nhiều màn hình nhưng luồng chính chưa chạy", "Đóng băng MVP; mọi nâng cấp đưa vào backlog sau nghiệm thu."),
        ("Sai doanh thu", "Client gửi số tiền hoặc giá cũ bị sửa", "Tính phí tại Spring Boot, transaction, lưu pricingRule và audit."),
        ("Tạo trùng lượt", "Bấm nhiều lần hoặc retry mạng", "Unique constraint, idempotencyKey và khóa nghiệp vụ."),
        ("Ghép hệ thống muộn", "Mỗi phần chạy riêng nhưng API lệch", "Swagger contract từ tuần 2; tích hợp hàng tuần."),
        ("Demo phụ thuộc camera", "Camera lỗi hoặc thiếu ánh sáng", "Có video/ảnh dự phòng và kịch bản demo offline.")
    ], [2100, 2800, 4460], 8.6)

    # Section 13
    doc.add_heading("13. HƯỚNG NÂNG CẤP SAU KHI HOÀN THÀNH", level=1)
    doc.add_heading("13.1. Nâng cấp AI", level=2)
    add_bullets(doc, [
        "Theo dõi biển số qua nhiều frame và voting kết quả.", "Huấn luyện CRNN riêng cho biển số Việt Nam.",
        "Nhận dạng loại xe, màu xe, hãng xe; so sánh ảnh xe vào/ra.", "Phát hiện biển số giả, che biển hoặc ảnh kém chất lượng.",
        "Active learning: dùng các trường hợp nhân viên sửa để bổ sung dataset."
    ])
    doc.add_heading("13.2. Nâng cấp vận hành", level=2)
    add_bullets(doc, [
        "Kết nối barie, RFID/NFC/QR và cảm biến vòng từ.", "Hỗ trợ nhiều cổng, nhiều làn và nhiều chi nhánh.",
        "Chế độ Desktop offline, hàng đợi sự kiện và đồng bộ khi có mạng.", "Theo dõi tình trạng camera và thiết bị theo thời gian thực."
    ])
    doc.add_heading("13.3. Nâng cấp Web Spring Boot", level=2)
    add_bullets(doc, [
        "Tách frontend React/Vue nhưng giữ Spring Boot REST API.", "WebSocket/SSE để cập nhật xe vào/ra trực tiếp trên dashboard.",
        "Redis cache và distributed lock khi mở rộng nhiều làn.", "Object storage S3/MinIO cho ảnh; chính sách lưu trữ và tự xóa.",
        "Thanh toán QR, hóa đơn điện tử, vé tháng online.", "SSO với tài khoản trường học/doanh nghiệp và phân quyền chi tiết."
    ])
    doc.add_heading("13.4. Hướng nghiên cứu nổi bật", level=2)
    add_callout(doc, "Đối chiếu đa phương thức", "Kết hợp biển số + màu xe + loại xe + ảnh toàn cảnh + mã thẻ để tạo điểm tin cậy tổng hợp. Hệ thống chỉ tự động cho qua khi điểm vượt ngưỡng; trường hợp còn lại chuyển cho nhân viên kiểm tra.")

    # Section 14
    doc.add_heading("14. CHECKLIST CHUẨN BỊ BẢO VỆ ĐỒ ÁN", level=1)
    doc.add_heading("14.1. Kịch bản demo 8-10 phút", level=2)
    add_numbers(doc, [
        "Đăng nhập Web bằng quản trị viên; giới thiệu dashboard.", "Tạo xe, gán thẻ và kiểm tra bảng giá.",
        "Desktop nhận diện xe đăng ký tại cổng vào và tạo lượt OPEN.", "Nhận diện một xe lạ hoặc biển số confidence thấp để minh họa cảnh báo.",
        "Cho xe đăng ký ra; xem phí dự kiến và xác nhận thanh toán.", "Quay lại Web tra cứu lịch sử, ảnh và doanh thu vừa phát sinh.",
        "Trình bày nhanh kết quả đánh giá YOLO/OCR và hướng nâng cấp."
    ])
    doc.add_heading("14.2. Tài liệu phải có", level=2)
    add_bullets(doc, [
        "Use case diagram, activity/sequence diagram, ERD và deployment diagram.", "Đặc tả API Swagger/OpenAPI.",
        "Mô tả dataset, cách chia train/validation/test và chỉ số ANPR.", "Test plan, test case và kết quả kiểm thử.",
        "Hướng dẫn cài đặt, dữ liệu mẫu, tài khoản demo và video dự phòng.", "Danh sách đóng góp của từng thành viên."
    ])
    doc.add_heading("14.3. Câu hỏi nhóm cần trả lời được", level=2)
    add_bullets(doc, [
        "Vì sao chọn Spring Boot + Thymeleaf thay vì React?", "Vì sao tách ANPR thành Python service?",
        "Hệ thống xử lý thế nào khi OCR sai hoặc mất mạng?", "Làm sao bảo đảm không tính phí hoặc tạo giao dịch hai lần?",
        "Số liệu độ chính xác được đo trên tập nào?", "Ảnh biển số và dữ liệu cá nhân được bảo vệ ra sao?"
    ])
    add_callout(doc, "Kết luận", "Sản phẩm thuyết phục nhất là sản phẩm có nghiệp vụ đúng, chạy ổn định và chứng minh được cách xử lý ngoại lệ. Độ chính xác AI cao là điểm cộng, nhưng phải đi cùng cơ chế xác nhận, lưu ảnh và truy vết.", fill="EAF5EF", border=GREEN)

    # Final handoff page
    add_page_break(doc)
    doc.add_heading("PHỤ LỤC A. VIỆC CẦN LÀM NGAY TRONG TUẦN ĐẦU", level=1)
    add_numbers(doc, [
        "Tạo repository với sáu thư mục: spring-web, desktop-wpf, anpr-service, database, docs, deployment.",
        "Tạo bảng Kanban gồm Backlog, Ready, In Progress, Review, Done.",
        "Vẽ ba sơ đồ đầu tiên: use case tổng quan, activity xe vào, activity xe ra.",
        "Chốt ERD phiên bản 1 và quy tắc: một thẻ - một xe; một xe - tối đa một lượt OPEN.",
        "Tạo Spring Boot project với Web, Thymeleaf, Security, Validation, JPA, PostgreSQL và Flyway.",
        "Tạo Swagger contract cho entry, exit-preview, exit-confirm và recognize.",
        "Tạo WPF skeleton mở webcam và Python FastAPI skeleton nhận file ảnh.",
        "Chuẩn bị 30-50 ảnh biển số đa dạng để chạy thử pipeline đầu tiên."
    ])
    doc.add_heading("PHỤ LỤC B. QUY ƯỚC GIT", level=1)
    add_bullets(doc, [
        "Nhánh: main, develop, feature/<module>-<feature>.", "Không commit trực tiếp vào main.",
        "Commit ngắn và có ý nghĩa, ví dụ: feat(session): add entry validation.", "Pull request phải mô tả cách test và ảnh giao diện nếu có.",
        "Không commit .env, mật khẩu, token, dataset lớn, model weight và ảnh người dùng thật."
    ])

    # Metadata
    doc.core_properties.title = "Kế hoạch đồ án bãi giữ xe ANPR - Java Spring Boot"
    doc.core_properties.subject = "Phân tích, thiết kế, phân công và lộ trình cho nhóm 3 người"
    doc.core_properties.author = "Nhóm đồ án"
    doc.core_properties.keywords = "ANPR, Spring Boot, Thymeleaf, YOLO, OCR, WPF, bãi giữ xe"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
