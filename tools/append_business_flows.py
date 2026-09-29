from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"D:\Kết quả tổng hợp mức phí khảo sát.docx")
OUTPUT = Path(r"D:\Code\KhoaLuan\Ket_qua_muc_phi_va_flow_nghiep_vu_hoan_chinh.docx")

FONT = "Times New Roman"
NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
LIGHT_YELLOW = "FFF2CC"
LIGHT_RED = "FCE4D6"
GREEN = "E2F0D9"
WHITE = "FFFFFF"
TEXT = "111111"


def set_run_font(run, size=12.5, bold=False, color=TEXT, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_inches):
    widths_dxa = [int(round(width * 1440)) for width in widths_inches]
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_inches[index])


def set_table_borders(table, color="B7C9DD", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def get_bullet_num_id(doc):
    cached = getattr(doc, "_flow_bullet_num_id", None)
    if cached is not None:
        return cached
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "420")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "420")
    ind.set(qn("w:hanging"), "210")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    doc._flow_bullet_num_id = num_id
    return num_id


def style_paragraph(paragraph, before=0, after=5, line=1.08, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.widow_control = True
    if keep:
        fmt.keep_with_next = True


def add_text(doc, text, bold=False, italic=False, color=TEXT, before=0, after=5, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent)
    style_paragraph(p, before=before, after=after)
    set_run_font(p.add_run(text), bold=bold, italic=italic, color=color)
    return p


def add_labeled(doc, label, text, color=TEXT):
    p = doc.add_paragraph()
    style_paragraph(p, after=4)
    set_run_font(p.add_run(label + ": "), bold=True, color=NAVY)
    set_run_font(p.add_run(text), color=color)
    return p


def add_heading(doc, text, level=1, page_break=False):
    p = doc.add_paragraph()
    if page_break:
        p.paragraph_format.page_break_before = True
    size = {1: 15, 2: 13.5, 3: 12.5}[level]
    color = NAVY if level < 3 else BLUE
    style_paragraph(p, before=9 if level > 1 else 0, after=6, keep=True)
    set_run_font(p.add_run(text), size=size, bold=True, color=color)
    return p


def add_summary(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, [6.25])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_margins(cell, top=120, bottom=120, start=150, end=150)
    p = cell.paragraphs[0]
    style_paragraph(p, after=0, line=1.05)
    set_run_font(p.add_run("Luồng tóm tắt: "), size=11.5, bold=True, color=NAVY)
    set_run_font(p.add_run(text), size=11.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_note(doc, label, text, fill=LIGHT_YELLOW):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.25])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=110, bottom=110, start=145, end=145)
    p = cell.paragraphs[0]
    style_paragraph(p, after=0)
    set_run_font(p.add_run(label + ": "), size=11.5, bold=True, color=NAVY)
    set_run_font(p.add_run(text), size=11.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_flow_table(doc, rows, headers=("Bước", "Xử lý nghiệp vụ", "Kết quả/Dữ liệu ghi nhận")):
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [0.65, 3.75, 1.85])
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0)
        set_run_font(p.add_run(text), size=10.5, bold=True, color=WHITE)
    for step, action, result in rows:
        cells = table.add_row().cells
        values = (str(step), action, result)
        for idx, value in enumerate(values):
            cell = cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if len(table.rows) % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, after=0, line=1.02)
            set_run_font(p.add_run(value), size=10.7, bold=(idx == 0), color=TEXT)
        prevent_row_split(table.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_rule_table(doc, rows, headers=("Tình huống/Điều kiện", "Quy tắc xử lý", "Trạng thái cuối")):
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [1.55, 3.55, 1.15])
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0)
        set_run_font(p.add_run(text), size=10.5, bold=True, color=WHITE)
    for condition, rule, status in rows:
        cells = table.add_row().cells
        for idx, value in enumerate((condition, rule, status)):
            cell = cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if len(table.rows) % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 2 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, after=0, line=1.02)
            set_run_font(p.add_run(value), size=10.5, bold=(idx == 2))
        prevent_row_split(table.rows[-1])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Paragraph")
    num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(get_bullet_num_id(doc)))
    num_pr.append(num_id)
    style_paragraph(p, after=3)
    set_run_font(p.add_run(text), size=12)
    return p


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def build():
    doc = Document(SOURCE)

    # Keep the existing FLOW title and replace only its four placeholder lines.
    for paragraph in list(doc.paragraphs[149:]):
        remove_paragraph(paragraph)

    add_text(
        doc,
        "Phần này mô tả luồng xử lý đề xuất cho hệ thống quản lý bãi đỗ xe chung cư, "
        "từ lúc tạo hồ sơ đến khi xe hoàn tất lượt ra/vào. Các ngưỡng AI và chính sách ngoại lệ "
        "phải được Ban quản lý cấu hình, kiểm thử và phê duyệt trước khi vận hành thực tế.",
        italic=True,
        color=NAVY,
        after=8,
    )

    add_heading(doc, "0. Phạm vi, tác nhân và nguyên tắc chung", level=1)
    add_labeled(doc, "Tác nhân", "Cư dân; thành viên hộ gia đình; khách vãng lai; nhân viên trạm gác; quản trị viên/Ban quản lý; hệ thống AI; cổng thanh toán.")
    add_labeled(doc, "Thiết bị", "Camera biển số; camera khuôn mặt; đầu đọc thẻ/QR; máy trạm WinForms tại cổng; máy chủ Spring Boot và cơ sở dữ liệu.")
    add_labeled(doc, "Đối tượng quản lý", "Hộ gia đình, cư dân, thành viên được ủy quyền, phương tiện, thẻ, gói gửi xe 30 ngày, lượt xe, ảnh bằng chứng, thanh toán và nhật ký xử lý.")
    add_labeled(doc, "Nguyên tắc an toàn", "AI chỉ đề xuất kết quả. Trường hợp gần ngưỡng, ảnh kém chất lượng hoặc dữ liệu mâu thuẫn phải chuyển nhân viên xác minh; barrier không tự mở khi chưa đủ điều kiện.")

    add_heading(doc, "0.1. Trạng thái chuẩn của dữ liệu", level=2)
    add_rule_table(
        doc,
        [
            ("Hồ sơ cư dân", "DRAFT -> PENDING_VERIFICATION -> ACTIVE; có thể chuyển SUSPENDED hoặc ARCHIVED.", "ACTIVE mới được sử dụng"),
            ("Thẻ gửi xe", "PENDING -> ACTIVE -> LOCKED/LOST/EXPIRED.", "Một thời điểm chỉ một trạng thái"),
            ("Gói 30 ngày", "PENDING_PAYMENT -> ACTIVE -> EXPIRING -> EXPIRED/CANCELLED.", "Kiểm tra tại thời điểm vào"),
            ("Lượt xe", "OPEN -> MANUAL_REVIEW/PAYMENT_PENDING -> CLOSED; CANCELLED chỉ dùng khi tạo nhầm.", "Không có hai lượt OPEN cho cùng xe"),
            ("Thanh toán", "UNPAID -> PAID; trường hợp nghiệp vụ đặc biệt có VOID hoặc REFUNDED kèm người duyệt.", "PAID mới đóng lượt có phí"),
        ],
    )

    add_heading(doc, "0.2. Mức quyết định tại cổng", level=2)
    add_rule_table(
        doc,
        [
            ("XANH", "Thẻ/biển số hợp lệ, chất lượng ảnh đạt, liveness đạt, khuôn mặt và quyền sử dụng xe hợp lệ.", "Cho phép"),
            ("VÀNG", "Điểm gần ngưỡng, ảnh bị lóa/mờ, CCCD cũ, nhận dạng biển số chưa chắc chắn hoặc dữ liệu cần xác minh.", "Kiểm tra thủ công"),
            ("ĐỎ", "Liveness thất bại, khuôn mặt không khớp rõ ràng, thẻ bị khóa, xe không được ủy quyền hoặc không tìm thấy lượt vào.", "Từ chối tự động"),
        ],
    )

    # RESIDENT FLOWS
    add_heading(doc, "1. FLOW NGHIỆP VỤ ĐỐI VỚI CƯ DÂN", level=1, page_break=True)

    add_heading(doc, "1.1. Đăng ký hộ gia đình, cư dân và gói gửi xe 30 ngày", level=2)
    add_summary(doc, "Tiếp nhận yêu cầu -> xác minh CCCD và khuôn mặt -> khai báo hộ gia đình -> đăng ký xe -> xác định mức phí -> thanh toán -> phát hành thẻ -> kích hoạt gói 30 ngày.")
    add_labeled(doc, "Điều kiện đầu vào", "Cư dân có thông tin căn hộ/hộ gia đình, CCCD hợp lệ và giấy tờ phương tiện cần đăng ký.")
    add_flow_table(
        doc,
        [
            (1, "Nhân viên tìm căn hộ/hộ gia đình. Nếu chưa tồn tại thì tạo hồ sơ hộ; nếu đã có thì kiểm tra trạng thái và người đại diện.", "Mã hộ, căn hộ, người đại diện"),
            (2, "Tiếp nhận CCCD của người đăng ký; nhập hoặc OCR họ tên, số định danh, ngày sinh và dữ liệu cần thiết. Nhân viên xác nhận lại dữ liệu OCR.", "Hồ sơ cư dân ở trạng thái PENDING_VERIFICATION"),
            (3, "Cắt riêng ảnh chân dung trên CCCD. Không sử dụng toàn bộ ảnh CCCD làm ảnh realtime.", "Ảnh tham chiếu CCCD"),
            (4, "Camera chụp ảnh đăng ký trong điều kiện đủ sáng, nhìn thẳng, không che mặt; thực hiện kiểm tra người thật (liveness).", "Ảnh đăng ký và kết quả liveness"),
            (5, "So sánh ảnh đăng ký với ảnh CCCD. Nếu gần ngưỡng thì chụp lại hoặc chuyển quản lý xác minh; nếu không khớp rõ ràng thì dừng đăng ký.", "Điểm đối chiếu CCCD - đăng ký"),
            (6, "Tạo mẫu đặc trưng khuôn mặt (face template/embedding), mã hóa và gắn với đúng cá nhân.", "Mẫu khuôn mặt bảo vệ bằng phân quyền"),
            (7, "Khai báo quan hệ của người này trong hộ: chủ hộ, vợ/chồng, con hoặc thành viên khác; xác định người được phép sử dụng từng xe.", "Danh sách thành viên và quyền sử dụng"),
            (8, "Nhập biển số, loại xe, số chỗ, dung tích xi-lanh hoặc loại động cơ theo giấy tờ xe. AI chỉ hỗ trợ phân loại thô; thông số chính xác phải lấy từ giấy tờ/nhập tay.", "Hồ sơ phương tiện"),
            (9, "Hệ thống ánh xạ phương tiện vào bảng giá cư dân: xe đạp/xe đạp điện; xe máy/xe điện thường; xe phân khối lớn; ô tô theo số chỗ.", "Mức phí 30 ngày trên từng xe"),
            (10, "Cư dân xác nhận danh sách xe, người được phép sử dụng và tổng tiền. Thu tiền theo phương thức được hỗ trợ.", "Phiếu thu PAID"),
            (11, "Phát hành thẻ/QR, gắn thẻ với xe và hồ sơ hộ. Mỗi xe được tính phí riêng theo bảng giá.", "Thẻ ACTIVE"),
            (12, "Kích hoạt thời hạn 30 ngày liên tiếp; ghi rõ thời điểm bắt đầu và kết thúc trên biên nhận.", "Gói ACTIVE và ngày hết hạn"),
        ],
    )
    add_note(doc, "Quy tắc bắt buộc", "Ba ảnh phải thuộc cùng người được đăng ký: ảnh CCCD, ảnh chụp đăng ký và ảnh realtime sau này. Không lấy ảnh của chủ xe để thay cho ảnh của thành viên khác trong gia đình.")

    add_heading(doc, "1.2. Đăng ký thêm thành viên được phép sử dụng xe", level=2)
    add_summary(doc, "Chọn hộ gia đình -> tạo/xác minh thành viên -> chụp ảnh đăng ký -> đối chiếu CCCD -> chọn xe được phép dùng -> lưu ủy quyền.")
    add_flow_table(
        doc,
        [
            (1, "Người đại diện hộ hoặc người có quyền gửi yêu cầu thêm thành viên.", "Yêu cầu PENDING"),
            (2, "Nhân viên kiểm tra quan hệ với hộ và thực hiện quy trình CCCD - ảnh đăng ký - liveness cho chính thành viên đó.", "Hồ sơ cá nhân đã xác minh"),
            (3, "Chọn từng phương tiện thành viên được phép sử dụng; có thể đặt ngày bắt đầu, ngày kết thúc hoặc quyền không thời hạn.", "Bản ghi ủy quyền theo người - xe"),
            (4, "Hệ thống kiểm tra trùng số định danh, trùng mẫu khuôn mặt và giới hạn chính sách của chung cư.", "Không tạo hồ sơ trùng"),
            (5, "Ban quản lý duyệt và kích hoạt. Từ thời điểm đó AI được phép so sánh realtime với mẫu của thành viên này.", "Ủy quyền ACTIVE"),
        ],
    )
    add_note(doc, "Ví dụ", "Xe đứng tên người chồng nhưng người vợ điều khiển: hệ thống tìm hộ gia đình từ xe/thẻ, xác định người vợ có ủy quyền, rồi so sánh ảnh realtime với ảnh CCCD và ảnh đăng ký của người vợ, không so sánh với ảnh người chồng.", fill=GREEN)

    add_heading(doc, "1.3. Xe cư dân đi vào", level=2)
    add_summary(doc, "Phát hiện xe -> đọc biển số và loại xe -> đọc thẻ -> kiểm tra gói 30 ngày -> xác định người có quyền -> chụp realtime/liveness -> đối chiếu -> mở barrier -> tạo lượt OPEN.")
    add_labeled(doc, "Điều kiện đầu vào", "Hồ sơ cư dân, xe, thẻ và gói gửi xe đã tồn tại; camera và máy trạm sẵn sàng.")
    add_flow_table(
        doc,
        [
            (1, "Cảm biến/camera phát hiện phương tiện tại làn vào và chụp nhiều khung hình.", "Tập khung hình đầu vào"),
            (2, "YOLO phát hiện vùng biển số; OCR đọc ký tự; chuẩn hóa bỏ khoảng trắng/dấu không cần thiết và trả độ tin cậy.", "Biển số dự đoán, ảnh cắt biển số"),
            (3, "Mô hình phân loại xác định nhóm xe thô: xe đạp, xe máy hoặc ô tô. Thông số phân khối/số chỗ lấy từ hồ sơ đăng ký, không suy đoán chỉ bằng ảnh.", "Loại xe quan sát"),
            (4, "Đọc thẻ/QR và tìm hồ sơ. Đối chiếu thẻ - biển số - loại xe - trạng thái xe.", "Xe/thẻ hợp lệ hoặc cảnh báo"),
            (5, "Kiểm tra gói 30 ngày có ACTIVE tại thời điểm vào và xe chưa có lượt OPEN.", "Điều kiện vào hợp lệ"),
            (6, "Lấy danh sách tất cả thành viên trong hộ đang được phép sử dụng xe này.", "Tập ứng viên giới hạn trong hộ"),
            (7, "Camera khuôn mặt chụp realtime; kiểm tra một khuôn mặt, độ sáng, độ nét, góc nhìn, che khuất và liveness.", "Ảnh realtime đạt chất lượng"),
            (8, "So sánh realtime với ảnh đăng ký/mẫu CCCD của người được phép. Ảnh đăng ký là nguồn chính; CCCD là nguồn đối chứng.", "Điểm khuôn mặt và người khớp"),
            (9, "Áp dụng mức XANH/VÀNG/ĐỎ. Mức VÀNG yêu cầu nhân viên kiểm tra và ghi lý do; mức ĐỎ không tự mở barrier.", "Quyết định có nhật ký"),
            (10, "Khi được phép, lưu thời gian vào, làn, nhân viên, biển số, loại xe, ảnh toàn cảnh, ảnh biển số, ảnh realtime và điểm AI.", "Lượt xe OPEN"),
            (11, "Mở barrier và hiển thị thông báo. Nếu barrier không phản hồi thì cảnh báo kỹ thuật, không tạo lượt OPEN thứ hai.", "Xe vào bãi an toàn"),
        ],
    )

    add_heading(doc, "1.4. Xe cư dân đi ra", level=2)
    add_summary(doc, "Đọc biển số/thẻ -> tìm lượt OPEN -> chụp ảnh ra -> so sánh xe và khuôn mặt -> kiểm tra quyền -> đóng lượt -> mở barrier.")
    add_flow_table(
        doc,
        [
            (1, "Chụp ảnh xe, biển số và khuôn mặt tại làn ra; thực hiện ANPR, phân loại xe và kiểm tra chất lượng ảnh.", "Dữ liệu xe ra"),
            (2, "Đọc thẻ/QR và tìm đúng lượt OPEN theo xe/thẻ. Không lấy một lượt CLOSED hoặc lượt của xe khác.", "Lượt vào tương ứng"),
            (3, "Đối chiếu biển số vào - ra, loại xe, đặc điểm phương tiện và thẻ.", "Kết quả khớp phương tiện"),
            (4, "Thực hiện liveness và so sánh khuôn mặt ra với hồ sơ người được phép; đồng thời dùng ảnh realtime lúc vào làm bằng chứng bổ sung.", "Kết quả khớp người"),
            (5, "Nếu người ra khác người vào nhưng đều thuộc danh sách được phép sử dụng xe, chuyển VÀNG để nhân viên xác nhận theo chính sách bàn giao xe.", "Quyết định bàn giao hợp lệ/không hợp lệ"),
            (6, "Gói tháng hết hạn trong khi xe đang ở trong bãi: cho phép xử lý ra an toàn, ghi cảnh báo và khóa lượt vào tiếp theo; không tự chuyển thành phí khách nếu chưa có chính sách/đồng ý.", "Cảnh báo gia hạn"),
            (7, "Khi đủ điều kiện, cập nhật thời gian ra, ảnh ra, điểm AI, người xử lý và trạng thái CLOSED.", "Lượt xe CLOSED"),
            (8, "Mở barrier; ghi nhận kết quả thiết bị. Nếu mở thủ công phải lưu tài khoản nhân viên và lý do.", "Hoàn tất xe ra"),
        ],
    )

    add_heading(doc, "1.5. Gia hạn, thay đổi và chấm dứt gói 30 ngày", level=2)
    add_flow_table(
        doc,
        [
            (1, "Trước ngày hết hạn theo số ngày cấu hình, hệ thống gửi/hiển thị cảnh báo cho cư dân và Ban quản lý.", "Gói EXPIRING"),
            (2, "Cư dân chọn xe cần gia hạn; hệ thống lấy mức phí đúng nhóm xe hiện tại.", "Đơn gia hạn"),
            (3, "Nếu gia hạn trước hạn, kỳ mới bắt đầu ngay sau thời điểm kết thúc kỳ hiện tại để không mất ngày.", "Thời hạn nối tiếp"),
            (4, "Nếu gia hạn sau hạn, kỳ mới bắt đầu tại thời điểm được kích hoạt hoặc ngày do Ban quản lý phê duyệt.", "Không hồi tố tự động"),
            (5, "Sau thanh toán thành công, cập nhật gói ACTIVE, biên nhận và doanh thu tháng.", "Gia hạn hoàn tất"),
            (6, "Khi đổi biển số/xe/thẻ, khóa liên kết cũ, xác minh giấy tờ mới, tính chênh lệch theo chính sách và lưu lịch sử thay đổi.", "Dữ liệu mới có hiệu lực"),
            (7, "Khi báo mất thẻ, chuyển thẻ LOST ngay; thẻ thay thế chỉ ACTIVE sau khi thẻ cũ bị vô hiệu hóa.", "Không tồn tại hai thẻ hợp lệ ngoài chính sách"),
        ],
    )
    add_note(doc, "Quy ước thời hạn đề xuất", "Một kỳ gồm 30 ngày liên tiếp. Hệ thống phải lưu cả thời điểm bắt đầu và thời điểm kết thúc đến giây; biên nhận phải hiển thị rõ hai mốc này để tránh hiểu nhầm là theo tháng dương lịch.")

    add_heading(doc, "1.6. Tính phí cư dân", level=2)
    add_rule_table(
        doc,
        [
            ("Xe đạp, xe đạp điện", "100.000 đồng/xe/30 ngày.", "Theo từng xe"),
            ("Xe máy, xe điện thông thường", "150.000 đồng/xe/30 ngày.", "Theo từng xe"),
            ("Xe máy phân khối lớn", "300.000 đồng/xe/30 ngày. Ngưỡng phân khối phải do Ban quản lý cấu hình và đối chiếu giấy tờ.", "Theo từng xe"),
            ("Ô tô 4-5 chỗ", "1.500.000 đồng/xe/30 ngày.", "Theo từng xe"),
            ("Ô tô 6-7 chỗ", "1.700.000 đồng/xe/30 ngày.", "Theo từng xe"),
            ("Ô tô 8-9 chỗ", "1.800.000 đồng/xe/30 ngày.", "Theo từng xe"),
        ],
        headers=("Nhóm phương tiện", "Mức phí và quy tắc", "Phạm vi"),
    )
    add_note(doc, "Ví dụ", "Hộ đăng ký 2 xe máy thông thường và 1 ô tô 4 chỗ: 2 x 150.000 + 1.500.000 = 1.800.000 đồng cho một kỳ 30 ngày.", fill=GREEN)

    # VISITOR FLOWS
    add_heading(doc, "2. FLOW NGHIỆP VỤ ĐỐI VỚI KHÁCH VÃNG LAI", level=1, page_break=True)

    add_heading(doc, "2.1. Khách vãng lai đi vào", level=2)
    add_summary(doc, "Phát hiện xe -> đọc biển số và loại xe -> chụp khuôn mặt realtime vào -> kiểm tra chất lượng/liveness -> phát hành vé/thẻ -> tạo lượt OPEN -> mở barrier.")
    add_flow_table(
        doc,
        [
            (1, "Camera chụp toàn cảnh, biển số và phương tiện; ANPR trả biển số cùng độ tin cậy.", "Dữ liệu xe vào"),
            (2, "Phân loại nhóm xe để chọn bảng giá. Nếu AI không chắc chắn, nhân viên chọn loại xe và hệ thống lưu thao tác thủ công.", "Nhóm tính phí"),
            (3, "Camera chụp khuôn mặt realtime của người điều khiển; kiểm tra chất lượng và liveness. Khách không phải cung cấp bộ ảnh CCCD - đăng ký như cư dân.", "Ảnh khuôn mặt vào"),
            (4, "Nếu biển số/ảnh mặt không đạt, chụp lại tối đa số lần cấu hình; vẫn thất bại thì nhân viên xác nhận thủ công trước khi cho vào.", "Dữ liệu đủ dùng hoặc MANUAL_REVIEW"),
            (5, "Phát hành vé/thẻ tạm, gắn duy nhất với biển số, loại xe, ảnh vào và thời gian vào.", "Vé khách ACTIVE"),
            (6, "Kiểm tra không có lượt OPEN trùng cho cùng biển số/thẻ.", "Không tạo lượt kép"),
            (7, "Lưu lượt OPEN và ảnh bằng chứng; barrier chỉ mở sau khi ghi dữ liệu thành công hoặc cơ chế offline đã cấp số lượt tạm.", "Lượt khách OPEN"),
        ],
    )

    add_heading(doc, "2.2. Khách vãng lai đi ra và thanh toán", level=2)
    add_summary(doc, "Đọc vé/thẻ và biển số -> tìm lượt OPEN -> chụp khuôn mặt ra -> so sánh vào/ra -> tính phí -> xác nhận thanh toán -> đóng lượt -> mở barrier.")
    add_flow_table(
        doc,
        [
            (1, "Đọc vé/thẻ, biển số và loại xe tại làn ra; tìm lượt OPEN tương ứng.", "Lượt gửi cần thanh toán"),
            (2, "So sánh biển số, loại xe và ảnh toàn cảnh vào - ra.", "Đối chiếu phương tiện"),
            (3, "Chụp khuôn mặt realtime ra, thực hiện liveness và so sánh với ảnh realtime lúc vào.", "Đối chiếu người gửi - người nhận"),
            (4, "Nếu kết quả khác người hoặc khác xe, chuyển MANUAL_REVIEW; không tự mở barrier chỉ dựa trên việc có vé.", "Kiểm tra an ninh"),
            (5, "Tính tổng thời gian từ thời điểm vào đến thời điểm ra theo thời gian máy chủ; chọn bảng giá đúng loại xe.", "Thời lượng và biểu phí"),
            (6, "Áp dụng công thức xe máy/xe đạp điện hoặc ô tô; lưu chi tiết từng thành phần phí để có thể giải thích cho khách.", "Số tiền phải thu"),
            (7, "Nhân viên/khách xác nhận số tiền và phương thức; hệ thống chỉ ghi PAID khi giao dịch thành công.", "Thanh toán PAID"),
            (8, "Đóng lượt, khóa vé/thẻ tạm để không tái sử dụng, lưu biên nhận và doanh thu theo ca.", "Lượt CLOSED"),
            (9, "Mở barrier và ghi kết quả thiết bị.", "Khách hoàn tất rời bãi"),
        ],
    )

    add_heading(doc, "2.3. Flow tính phí xe máy vãng lai", level=2)
    add_summary(doc, "Tính thời lượng -> kiểm tra qua mốc 00:00/chu kỳ 24 giờ -> nếu qua đêm tính theo chu kỳ; nếu không, xác định ban ngày hay ban đêm -> trả phí.")
    add_flow_table(
        doc,
        [
            (1, "Tính thời lượng thực tế = thời điểm ra - thời điểm vào. Nếu thời điểm ra nhỏ hơn vào hoặc dữ liệu thiếu thì dừng và báo lỗi.", "Tổng số phút/giờ"),
            (2, "Kiểm tra lượt có đi qua mốc 00:00 hoặc kéo dài từ 24 giờ trở lên hay không.", "Nhánh qua đêm/không qua đêm"),
            (3, "Nếu qua đêm: số chu kỳ = làm tròn lên (tổng thời gian / 24 giờ). Đúng 24 giờ là 1 chu kỳ; 24 giờ 1 phút là 2 chu kỳ.", "Số chu kỳ"),
            (4, "Phí qua đêm xe máy = số chu kỳ x 10.000 đồng.", "Phí chu kỳ"),
            (5, "Nếu không qua đêm và toàn bộ lượt nằm trong 06:00-17:59: áp dụng 5.000 đồng/lượt ban ngày.", "Phí ban ngày"),
            (6, "Nếu không qua ngày mới nhưng lượt có thời gian từ 18:00-05:59: áp dụng 8.000 đồng/lượt ban đêm.", "Phí ban đêm"),
            (7, "Lưu loại công thức, thời lượng, chu kỳ và số tiền cuối cùng vào chi tiết thanh toán.", "Kết quả có thể kiểm tra lại"),
        ],
    )
    add_note(doc, "Ví dụ", "Xe máy vào 08:00 ngày 10/09 và ra 10:00 ngày 11/09: 26 giờ; số chu kỳ = làm tròn lên 26/24 = 2; phí = 2 x 10.000 = 20.000 đồng.", fill=GREEN)

    add_heading(doc, "2.4. Flow tính phí xe đạp và xe đạp điện vãng lai", level=2)
    add_rule_table(
        doc,
        [
            ("Không qua đêm, chỉ trong 06:00-17:59", "Một lượt ban ngày.", "2.000 đồng"),
            ("Không qua ngày mới, có thời gian 18:00-05:59", "Một lượt ban đêm.", "3.000 đồng"),
            ("Qua 00:00 hoặc từ 24 giờ", "Số chu kỳ = làm tròn lên tổng thời gian/24 giờ.", "5.000 đồng x chu kỳ"),
        ],
        headers=("Điều kiện", "Cách tính", "Phí"),
    )

    add_heading(doc, "2.5. Flow tính phí ô tô vãng lai", level=2)
    add_summary(doc, "Tính phí thời gian -> đếm số mốc 00:00 -> tính phí tối thiểu qua đêm -> lấy mức cao hơn -> lưu chi tiết.")
    add_flow_table(
        doc,
        [
            (1, "Tính tổng thời gian gửi theo giờ và phút.", "Tổng thời lượng"),
            (2, "Nếu không quá 4 giờ, phí thời gian = 40.000 đồng.", "Phí cơ bản"),
            (3, "Nếu trên 4 giờ, thời gian vượt = tổng thời gian - 4 giờ; số lần tính thêm = làm tròn lên thời gian vượt/2 giờ.", "Số lần 2 giờ"),
            (4, "Phí thời gian = 40.000 + số lần tính thêm x 20.000 đồng.", "Phí theo thời gian"),
            (5, "Đếm số mốc 00:00 mà lượt xe đi qua. Vào ngày 10, ra ngày 12 là 2 mốc.", "Số đêm"),
            (6, "Phí tối thiểu qua đêm = số đêm x 100.000 đồng.", "Phí tối thiểu"),
            (7, "Phí cuối cùng = mức cao hơn giữa phí theo thời gian và phí tối thiểu qua đêm.", "Số tiền phải thu"),
            (8, "Lưu cả hai kết quả, số đêm và giá trị được chọn để nhân viên giải thích khi cần.", "Chi tiết tính phí"),
        ],
    )
    add_rule_table(
        doc,
        [
            ("23:00-02:00 hôm sau", "Phí thời gian 40.000; 1 đêm x 100.000.", "100.000 đồng"),
            ("20:00-07:00 hôm sau", "11 giờ: phí thời gian 120.000; tối thiểu 1 đêm 100.000.", "120.000 đồng"),
            ("08:00 ngày 10-10:00 ngày 11", "26 giờ: phí thời gian 260.000; tối thiểu 1 đêm 100.000.", "260.000 đồng"),
            ("20:00 ngày 10-07:00 ngày 12", "35 giờ: phí thời gian 360.000; tối thiểu 2 đêm 200.000.", "360.000 đồng"),
        ],
        headers=("Ví dụ", "So sánh", "Phí cuối"),
    )

    # COMMON AI AND EXCEPTIONS
    add_heading(doc, "3. FLOW AI DÙNG CHUNG CHO CƯ DÂN VÀ KHÁCH", level=1, page_break=True)

    add_heading(doc, "3.1. Nhận dạng biển số và phân loại phương tiện", level=2)
    add_flow_table(
        doc,
        [
            (1, "Chọn khung hình tốt nhất từ camera dựa trên độ nét, độ sáng và mức che khuất.", "Khung hình chuẩn"),
            (2, "YOLO phát hiện và khoanh vùng biển số; lưu tọa độ và độ tin cậy.", "Vùng biển số"),
            (3, "Tiền xử lý vùng biển số: cân bằng sáng, giảm nhiễu, hiệu chỉnh nghiêng khi cần.", "Ảnh biển số đã chuẩn hóa"),
            (4, "OCR đọc ký tự; chuẩn hóa định dạng; so sánh nhiều khung hình để chọn kết quả ổn định.", "Chuỗi biển số và điểm OCR"),
            (5, "Mô hình phát hiện phương tiện trả nhóm xe quan sát. Đối với cư dân, đối chiếu thêm nhóm xe đã đăng ký.", "Loại xe AI"),
            (6, "Nếu biển số hoặc loại xe mâu thuẫn dữ liệu, không tự sửa hồ sơ; chuyển nhân viên xác nhận và lưu cả kết quả AI lẫn kết quả thủ công.", "Kết quả cuối có nguồn gốc"),
        ],
    )
    add_note(doc, "Giới hạn kỹ thuật", "Không nên dùng hình ảnh để kết luận chính xác dung tích xi-lanh, số chỗ hoặc loại động cơ điện/xăng trong mọi trường hợp. Các thuộc tính quyết định mức phí cư dân phải lấy từ giấy tờ xe hoặc dữ liệu đã được nhân viên xác minh.")

    add_heading(doc, "3.2. Đối chiếu ba ảnh khuôn mặt của cư dân", level=2)
    add_summary(doc, "Lúc đăng ký: CCCD <-> ảnh đăng ký. Tại cổng: realtime <-> ảnh đăng ký (chính) và realtime <-> mẫu CCCD (đối chứng). Khi ra: bổ sung realtime ra <-> realtime vào.")
    add_flow_table(
        doc,
        [
            (1, "Phát hiện khuôn mặt và kiểm tra chỉ có đúng người điều khiển trong vùng xác minh.", "Vùng khuôn mặt"),
            (2, "Đánh giá chất lượng: độ sáng, cháy sáng, độ nét, góc quay, kích thước khuôn mặt và mức che khuất.", "PASS hoặc yêu cầu chụp lại"),
            (3, "Thực hiện liveness/anti-spoofing để hạn chế ảnh in, ảnh điện thoại hoặc video giả.", "Liveness PASS/FAIL"),
            (4, "Căn chỉnh khuôn mặt theo các điểm mắt, mũi, miệng; tạo vector đặc trưng bằng mô hình chuyên dụng.", "Face embedding"),
            (5, "So sánh realtime với ảnh đăng ký của người có quyền sử dụng xe; đây là phép so sánh chính tại cổng.", "Điểm chính"),
            (6, "So sánh realtime với mẫu tạo từ ảnh CCCD để đối chứng; CCCD cũ/chất lượng thấp có thể chuyển VÀNG thay vì loại ngay.", "Điểm hỗ trợ"),
            (7, "Áp dụng ngưỡng đã hiệu chỉnh trên dữ liệu kiểm thử thực tế; không hiểu điểm cosine là phần trăm xác suất.", "MATCH/REVIEW/NO_MATCH"),
            (8, "Lưu phiên bản mô hình, ngưỡng, điểm, cảnh báo chất lượng và người xử lý để phục vụ kiểm tra sau này.", "Nhật ký AI"),
        ],
    )

    add_heading(doc, "3.3. Xử lý ảnh cháy sáng, tối, mờ hoặc bị che", level=2)
    add_rule_table(
        doc,
        [
            ("Cháy sáng/đèn pha", "Bật WDR/HDR, khóa phơi sáng vùng mặt/biển số, chọn khung khác; yêu cầu xe dừng đúng vạch.", "Chụp lại"),
            ("Ảnh quá tối", "Bật đèn hỗ trợ/hồng ngoại, tăng sáng có kiểm soát và chọn khung nét nhất.", "Chụp lại"),
            ("Mờ do chuyển động", "Tăng tốc màn trập, chụp chuỗi khung, yêu cầu dừng xe.", "Chụp lại"),
            ("Mũ bảo hiểm/kính tối/khẩu trang", "Yêu cầu lộ đủ vùng mặt theo chính sách và thực hiện liveness; không ép AI nhận dạng khi thiếu dữ liệu.", "REVIEW nếu vẫn che"),
            ("Biển số bẩn/bị che", "Chụp camera góc phụ; nhân viên nhập tay và lưu ảnh/lý do.", "MANUAL_REVIEW"),
            ("Sau số lần chụp lại tối đa", "Không lặp vô hạn; chuyển nhân viên và giữ barrier đóng cho đến khi có quyết định.", "MANUAL_REVIEW"),
        ],
    )

    # EXCEPTIONS
    add_heading(doc, "4. CÁC TRƯỜNG HỢP NGOẠI LỆ VÀ HƯỚNG XỬ LÝ", level=1, page_break=True)
    add_heading(doc, "4.1. Ngoại lệ đối với cư dân", level=2)
    add_rule_table(
        doc,
        [
            ("Quên/mất/hỏng thẻ", "Tra cứu biển số và hộ; xác minh khuôn mặt/liveness; thẻ mất phải chuyển LOST; mở thủ công cần ghi lý do.", "REVIEW"),
            ("Gói 30 ngày hết hạn trước khi vào", "Thông báo gia hạn. Chỉ tạo lượt khách khi cư dân đồng ý và chính sách cho phép; không âm thầm đổi loại phí.", "Từ chối hoặc tạo lượt khách"),
            ("Gói hết hạn khi xe đang trong bãi", "Cho xử lý ra an toàn, cảnh báo và khóa lượt vào kế tiếp; không hồi tố phí khách nếu chính sách chưa quy định.", "CLOSED + cảnh báo"),
            ("Người thân chưa được ủy quyền", "Không dùng ảnh chủ xe để thay thế. Liên hệ người đại diện/ban quản lý và tạo phê duyệt tạm nếu chính sách cho phép.", "REVIEW/REJECT"),
            ("Người vào và người ra khác nhau", "Nếu cả hai đều được ủy quyền thì nhân viên xác nhận bàn giao; nếu không thì giữ xe và báo quản lý.", "REVIEW"),
            ("Biển số khác hồ sơ", "Chụp lại; kiểm tra thay biển/xe mới; không tự cập nhật biển số từ kết quả OCR.", "REVIEW/REJECT"),
            ("Loại xe AI khác hồ sơ", "Nhân viên kiểm tra giấy tờ/quan sát; lưu kết quả thủ công và tạo yêu cầu cập nhật nếu cần.", "REVIEW"),
            ("Có lượt OPEN trùng", "Không tạo lượt mới; tìm lượt trước, kiểm tra thao tác cổng và chỉ hủy/sửa khi có quyền.", "Giải quyết dữ liệu"),
            ("Liveness thất bại", "Chụp lại với hướng dẫn; nếu tiếp tục thất bại thì không tự mở barrier.", "REJECT/REVIEW"),
            ("Camera/mạng/máy chủ lỗi", "Chuyển chế độ offline có số lượt tạm, chụp ảnh cục bộ và đồng bộ sau; mở thủ công phải có nhật ký.", "OFFLINE_MODE"),
        ],
        headers=("Ngoại lệ", "Hướng xử lý", "Kết quả"),
    )

    add_heading(doc, "4.2. Ngoại lệ đối với khách vãng lai", level=2)
    add_rule_table(
        doc,
        [
            ("Mất vé/thẻ", "Tra biển số, ảnh xe, khuôn mặt vào/ra và thời gian; yêu cầu quản lý duyệt. Phí mất vé chỉ thu khi bảng chính sách có cấu hình.", "MANUAL_REVIEW"),
            ("Vé đúng nhưng biển số khác", "Không cho ra tự động; kiểm tra ảnh vào, giấy tờ và camera. Lưu biên bản xử lý.", "REJECT/REVIEW"),
            ("Khuôn mặt vào/ra không khớp", "Xác minh trường hợp người khác lấy xe; kiểm tra quyền sở hữu/ủy quyền theo quy định và cần quản lý duyệt.", "MANUAL_REVIEW"),
            ("Không tìm thấy lượt OPEN", "Tìm theo biển số chuẩn hóa, thẻ, khung giờ và làn; không tạo phí ước lượng khi chưa có bằng chứng.", "MANUAL_REVIEW"),
            ("AI không xác định loại xe", "Nhân viên chọn loại xe trước khi tính phí và lưu nguồn MANUAL.", "Tiếp tục tính phí"),
            ("Khách không đủ tiền/lỗi thanh toán", "Giữ PAYMENT_PENDING; áp dụng quy trình xử lý của Ban quản lý, không đánh dấu PAID giả.", "PAYMENT_PENDING"),
            ("Tranh chấp thời gian/phí", "Hiển thị ảnh vào/ra, mốc thời gian và chi tiết công thức; điều chỉnh chỉ bởi người có quyền và phải có lý do.", "PAID/ADJUSTED"),
            ("Barrier mở nhưng chưa đóng lượt", "Cảnh báo ngay; nhân viên xác nhận xe đã ra và đóng lượt có nhật ký, tránh tính phí tiếp.", "CLOSED_MANUAL"),
        ],
        headers=("Ngoại lệ", "Hướng xử lý", "Kết quả"),
    )

    # PAYMENT, REPORTING, OFFLINE
    add_heading(doc, "5. THANH TOÁN, DOANH THU VÀ ĐỒNG BỘ DỮ LIỆU", level=1, page_break=True)
    add_heading(doc, "5.1. Flow thanh toán và ghi nhận doanh thu", level=2)
    add_flow_table(
        doc,
        [
            (1, "Tạo khoản phải thu từ gói 30 ngày hoặc lượt khách; lưu bảng giá/phiên bản công thức được áp dụng.", "Invoice/charge UNPAID"),
            (2, "Hiển thị số tiền, chi tiết tính và phương thức thanh toán cho người nộp.", "Xác nhận trước thanh toán"),
            (3, "Nhận kết quả tiền mặt/chuyển khoản/QR. Giao dịch điện tử phải có mã đối soát; tiền mặt phải gắn ca nhân viên.", "Payment result"),
            (4, "Chỉ khi thành công mới chuyển PAID, phát hành biên nhận và kích hoạt gói/cho đóng lượt.", "Biên nhận và doanh thu"),
            (5, "Nếu hoàn/hủy/điều chỉnh, yêu cầu quyền quản lý, lý do và liên kết giao dịch gốc.", "REFUNDED/VOID/ADJUSTED"),
            (6, "Cuối ca, tổng hợp tiền mặt, điện tử, lượt miễn/điều chỉnh và chênh lệch; nhân viên bàn giao, quản lý xác nhận.", "Báo cáo ca"),
        ],
    )

    add_heading(doc, "5.2. Chỉ tiêu báo cáo", level=2)
    add_bullet(doc, "Doanh thu theo ngày, tháng, ca làm việc, loại xe, loại khách và phương thức thanh toán.")
    add_bullet(doc, "Số gói 30 ngày mới, gia hạn, hết hạn và doanh thu theo nhóm phương tiện.")
    add_bullet(doc, "Số lượt vào/ra, số xe đang trong bãi, thời gian gửi trung bình và lượt quá 24 giờ.")
    add_bullet(doc, "Tỷ lệ nhận dạng biển số thành công, tỷ lệ ảnh khuôn mặt đạt chất lượng, số ca REVIEW/REJECT và nguyên nhân.")
    add_bullet(doc, "Danh sách giao dịch điều chỉnh, mở barrier thủ công, mất vé và chênh lệch cuối ca.")

    add_heading(doc, "5.3. Flow khi mất kết nối", level=2)
    add_flow_table(
        doc,
        [
            (1, "Máy trạm phát hiện không kết nối được API/cơ sở dữ liệu trong thời gian cấu hình.", "Chuyển OFFLINE_MODE"),
            (2, "Tải bộ dữ liệu tối thiểu đã mã hóa: thẻ/xe cư dân ACTIVE, cấu hình phí và dải số lượt tạm.", "Có thể vận hành hạn chế"),
            (3, "Lưu cục bộ ảnh, thời gian, biển số, thẻ, thao tác nhân viên và mã lượt tạm; không xóa khi chưa đồng bộ.", "Hàng đợi offline"),
            (4, "Các trường hợp không đủ dữ liệu phải kiểm tra thủ công; không giả định là hợp lệ chỉ vì hệ thống mất mạng.", "Quyết định có trách nhiệm"),
            (5, "Khi có mạng, gửi theo thứ tự thời gian với khóa chống trùng; máy chủ trả mã lượt chính thức.", "Đồng bộ idempotent"),
            (6, "Nếu xung đột lượt OPEN/thanh toán, giữ trạng thái SYNC_CONFLICT để quản lý xử lý, không tự ghi đè.", "Danh sách xung đột"),
        ],
    )

    add_heading(doc, "5.4. Bảo vệ dữ liệu và nhật ký", level=2)
    add_bullet(doc, "Phân quyền tối thiểu: nhân viên trạm gác chỉ xem dữ liệu cần cho ca trực; Ban quản lý mới được sửa hồ sơ và duyệt ngoại lệ.")
    add_bullet(doc, "Mã hóa dữ liệu nhạy cảm khi lưu và truyền; ưu tiên lưu face template thay vì sử dụng ảnh thô cho mọi phép so sánh.")
    add_bullet(doc, "Mọi lần xem/sửa/xóa hồ sơ, mở barrier thủ công, thay đổi phí và điều chỉnh thanh toán phải có audit log.")
    add_bullet(doc, "Quy định rõ thời hạn lưu ảnh vào/ra, ảnh CCCD, ảnh đăng ký và nhật ký; xóa/ẩn danh khi hết mục đích theo chính sách được phê duyệt.")
    add_bullet(doc, "Dữ liệu demo phải dùng người tự nguyện hoặc dữ liệu giả lập; không gửi ảnh CCCD/khuôn mặt thật lên dịch vụ AI công cộng khi chưa có cơ sở và chấp thuận phù hợp.")

    # SCENARIOS
    add_heading(doc, "6. CÁC KỊCH BẢN KIỂM THỬ NGHIỆP VỤ ĐẦU CUỐI", level=1, page_break=True)
    add_rule_table(
        doc,
        [
            ("Cư dân chính chủ vào/ra", "Thẻ, biển số, gói tháng, liveness và khuôn mặt đều hợp lệ; tạo OPEN khi vào và CLOSED khi ra.", "Cho phép"),
            ("Vợ dùng xe đứng tên chồng", "Vợ thuộc hộ và có ủy quyền; so sánh realtime với ảnh của vợ; không dùng ảnh chồng.", "Cho phép/REVIEW theo chính sách"),
            ("Người thân chưa đăng ký", "Biển số/thẻ đúng nhưng không có hồ sơ khuôn mặt/ủy quyền của người lái.", "Không tự động cho phép"),
            ("Gói tháng vừa hết hạn", "Cảnh báo gia hạn; không tự coi là khách khi chưa có lựa chọn nghiệp vụ.", "Gia hạn hoặc tạo lượt khách"),
            ("Khách xe máy gửi 26 giờ", "2 chu kỳ x 10.000 đồng.", "20.000 đồng"),
            ("Khách ô tô 23:00-02:00", "Phí thời gian 40.000; tối thiểu qua 1 đêm 100.000; lấy mức cao hơn.", "100.000 đồng"),
            ("Ảnh mặt bị cháy sáng", "Không so sánh ngay; điều chỉnh camera/chụp lại; quá số lần thì nhân viên xử lý.", "REVIEW"),
            ("Biển số OCR sai một ký tự", "Đối chiếu nhiều khung, thẻ và hồ sơ; nhân viên sửa có lưu kết quả AI ban đầu.", "Cho phép sau xác minh"),
            ("Khách mất vé", "Tra ảnh/biển số/khuôn mặt vào-ra; quản lý duyệt và áp chính sách mất vé nếu được cấu hình.", "MANUAL_REVIEW"),
            ("Mất mạng khi xe vào", "Tạo số lượt tạm và lưu cục bộ; đồng bộ chống trùng khi có mạng.", "OFFLINE_OPEN"),
        ],
        headers=("Kịch bản", "Kết quả mong đợi", "Trạng thái"),
    )

    add_heading(doc, "6.1. Tiêu chí nghiệm thu tối thiểu", level=2)
    add_bullet(doc, "Mỗi lượt xe có thể truy ngược đầy đủ dữ liệu vào, dữ liệu ra, quyết định AI, quyết định nhân viên và thanh toán.")
    add_bullet(doc, "Không tạo hai lượt OPEN cho cùng một xe nếu chưa xử lý lượt trước.")
    add_bullet(doc, "Không tự mở barrier khi liveness thất bại, không tìm thấy lượt vào hoặc dữ liệu xe/người mâu thuẫn rõ ràng.")
    add_bullet(doc, "Công thức phí cho kết quả đúng ở các mốc biên: 4 giờ, 4 giờ 1 phút, đúng 24 giờ, 24 giờ 1 phút và nhiều mốc 00:00.")
    add_bullet(doc, "Thành viên hộ gia đình chỉ được xác minh bằng hồ sơ của chính họ và chỉ được dùng xe đã được ủy quyền.")
    add_bullet(doc, "Mọi xử lý thủ công đều có tài khoản, thời gian, lý do và ảnh bằng chứng.")
    add_bullet(doc, "Hệ thống khôi phục sau mất mạng mà không làm trùng lượt, trùng thanh toán hoặc mất ảnh bằng chứng.")

    # Add a simple footer page number field to the whole document without changing the body design.
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, after=0)
    run = p.add_run("Trang ")
    set_run_font(run, size=10, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)

    doc.core_properties.title = "Kết quả tổng hợp mức phí khảo sát và flow nghiệp vụ"
    doc.core_properties.subject = "Quản lý bãi đỗ xe chung cư tích hợp ANPR và đối chiếu khuôn mặt"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
